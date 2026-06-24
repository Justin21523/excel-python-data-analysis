from pathlib import Path

from openpyxl import load_workbook

from excel_analysis.data import create_retail_workbook, generate_retail_data
from excel_analysis.io import ExcelReader
from excel_analysis.profiling import DataProfiler
from excel_analysis.services import AnalysisPipeline


def test_generate_retail_data_has_required_tables():
    data = generate_retail_data()
    assert {"orders", "order_items", "customers", "products", "inventory"} <= set(data)
    assert len(data["orders"]) > 100


def test_excel_reader_reads_multiple_sheets(tmp_path):
    workbook = create_retail_workbook(tmp_path / "demo.xlsx")
    sheets = ExcelReader().read(workbook)
    assert "orders" in sheets
    assert "order_items" in sheets


def test_profiler_detects_quality_issues():
    data = generate_retail_data()
    profile = DataProfiler().profile("orders", data["orders"])
    assert profile.rows == len(data["orders"])
    assert profile.quality_score < 100
    assert any(issue.issue_type == "missing_values" for issue in profile.issues)


def test_pipeline_creates_excel_report(tmp_path):
    workbook = create_retail_workbook(tmp_path / "demo.xlsx")
    result = AnalysisPipeline().run(workbook, tmp_path)
    report_path = result.artifacts[0].path
    assert report_path.exists()
    wb = load_workbook(report_path)
    expected = {"Summary", "Data Quality", "Sales", "Customers", "Inventory", "Market Basket", "Cleaned Sample", "Metadata"}
    assert expected <= set(wb.sheetnames)
    assert wb["Sales"].freeze_panes == "A2"

