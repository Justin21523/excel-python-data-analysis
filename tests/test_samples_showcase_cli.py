from openpyxl import load_workbook

from excel_analysis.cli import main
from excel_analysis.data import (
    create_all_sample_workbooks,
    create_customer_segmentation_workbook,
    create_finance_workbook,
    create_inventory_planning_workbook,
    create_messy_sales_workbook,
)
from excel_analysis.reporting import OpenpyxlFeatureShowcase


def test_sample_business_workbooks_are_generated(tmp_path):
    paths = [
        create_messy_sales_workbook(tmp_path / "messy.xlsx"),
        create_finance_workbook(tmp_path / "finance.xlsx"),
        create_inventory_planning_workbook(tmp_path / "inventory.xlsx"),
        create_customer_segmentation_workbook(tmp_path / "customer.xlsx"),
    ]

    assert all(path.exists() for path in paths)
    assert "messy_sales" in load_workbook(paths[0]).sheetnames
    assert "cash_flow" in load_workbook(paths[1]).sheetnames
    assert "inventory_plan" in load_workbook(paths[2]).sheetnames
    assert "engagement" in load_workbook(paths[3]).sheetnames


def test_create_all_sample_workbooks(tmp_path):
    paths = create_all_sample_workbooks(tmp_path)

    assert len(paths) == 5
    assert all(path.exists() for path in paths)


def test_openpyxl_showcase_contains_required_features(tmp_path):
    artifact = OpenpyxlFeatureShowcase().build(tmp_path / "showcase.xlsx")
    wb = load_workbook(artifact.path, data_only=False)

    expected = {
        "Capabilities",
        "Formula Demo",
        "Named Styles",
        "Comments Links",
        "Data Validation",
        "Protection",
        "Conditional Formatting",
        "Charts",
    }
    assert expected <= set(wb.sheetnames)
    assert wb["Formula Demo"]["D2"].value.startswith("=")
    assert wb["Comments Links"]["A1"].comment is not None
    assert len(wb["Data Validation"].data_validations.dataValidation) >= 3
    assert wb["Protection"].protection.sheet is True
    assert len(wb["Conditional Formatting"].conditional_formatting) >= 3
    assert len(wb["Charts"]._charts) == 3


def test_cli_commands_generate_clean_and_export(tmp_path):
    sample_dir = tmp_path / "samples"
    report_path = tmp_path / "report.xlsx"
    showcase_path = tmp_path / "showcase.xlsx"
    cleaned_path = tmp_path / "cleaned.xlsx"

    assert main(["generate-sample", "--case", "all", "--output-dir", str(sample_dir)]) == 0
    assert main(["clean", "--input", str(sample_dir / "messy_sales.xlsx"), "--output", str(cleaned_path)]) == 0
    assert main(["profile", "--input", str(sample_dir / "retail_demo.xlsx")]) == 0
    assert main(["export-report", "--input", str(sample_dir / "retail_demo.xlsx"), "--output", str(report_path)]) == 0
    assert main(["showcase-openpyxl", "--output", str(showcase_path)]) == 0

    assert cleaned_path.exists()
    assert report_path.exists()
    assert showcase_path.exists()

