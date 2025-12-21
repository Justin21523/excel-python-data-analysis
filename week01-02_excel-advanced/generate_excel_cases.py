#!/usr/bin/env python3
"""
生成 Week 1-2 Excel 練習檔案
包含完整的範例資料和公式
"""

import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from faker import Faker

# 設定中文 Faker
fake = Faker("zh_TW")
np.random.seed(42)

# 輸出目錄
OUTPUT_DIR = Path(
    "/home/justin/web-projects/excel-python-data-analysis/week01-02_excel-advanced"
)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

print("=" * 80)
print("🚀 開始生成 Excel 練習檔案...")
print("=" * 80)


# =============================================================================
# Case 01: 動態銷售報表
# =============================================================================
def generate_case01():
    print("\n📝 生成 Case 01: 動態銷售報表...")

    # 生成訂單資料（6 個月）
    n_orders = 1000

    start_date = datetime(2024, 6, 1)
    end_date = datetime(2024, 11, 30)

    # 產品資料
    categories = [
        "電腦周邊",
        "手機配件",
        "家電",
        "服飾",
        "美妝",
        "食品",
        "書籍",
        "運動用品",
    ]
    products = {
        "電腦周邊": ["無線滑鼠", "機械鍵盤", "USB-C Hub", "筆電支架", "外接硬碟"],
        "手機配件": ["手機殼", "藍牙耳機", "充電線", "行動電源", "螢幕保護貼"],
        "家電": ["吸塵器", "咖啡機", "電風扇", "空氣清淨機", "烤箱"],
        "服飾": ["T恤", "牛仔褲", "外套", "運動鞋", "背包"],
        "美妝": ["面膜", "精華液", "洗面乳", "化妝水", "護髮素"],
        "食品": ["即溶咖啡", "餅乾", "泡麵", "零食", "飲料"],
        "書籍": ["商業書籍", "小說", "漫畫", "雜誌", "童書"],
        "運動用品": ["瑜珈墊", "啞鈴", "運動服", "跑步鞋", "水壺"],
    }

    # 價格範圍
    price_ranges = {
        "電腦周邊": (300, 3000),
        "手機配件": (100, 2000),
        "家電": (1000, 8000),
        "服飾": (500, 3000),
        "美妝": (200, 1500),
        "食品": (50, 500),
        "書籍": (150, 800),
        "運動用品": (300, 2500),
    }

    orders = []
    for i in range(n_orders):
        order_date = start_date + timedelta(
            days=np.random.randint(0, (end_date - start_date).days)
        )

        category = np.random.choice(categories)
        product = np.random.choice(products[category])
        price_min, price_max = price_ranges[category]
        price = np.random.randint(price_min, price_max)
        quantity = np.random.choice(
            [1, 1, 1, 2, 2, 3], p=[0.5, 0.2, 0.1, 0.1, 0.05, 0.05]
        )

        orders.append(
            {
                "訂單編號": f"ORD{i+1:05d}",
                "訂單日期": order_date,
                "產品類別": category,
                "產品名稱": product,
                "單價": price,
                "數量": quantity,
                "金額": price * quantity,
                "客戶姓名": fake.name(),
                "地區": np.random.choice(
                    ["北部", "中部", "南部", "東部"], p=[0.4, 0.3, 0.2, 0.1]
                ),
                "訂單狀態": np.random.choice(
                    ["已完成", "已完成", "已完成", "處理中", "已取消"],
                    p=[0.7, 0.15, 0.05, 0.08, 0.02],
                ),
            }
        )

    df_orders = pd.DataFrame(orders)

    # 產品主檔
    product_master = []
    for cat, prods in products.items():
        for prod in prods:
            price_min, price_max = price_ranges[cat]
            product_master.append(
                {
                    "產品編號": f"P{len(product_master)+1:04d}",
                    "產品類別": cat,
                    "產品名稱": prod,
                    "建議售價": np.random.randint(price_min, price_max),
                    "供應商": fake.company(),
                    "庫存數量": np.random.randint(10, 500),
                }
            )

    df_products = pd.DataFrame(product_master)

    # 創建 Excel
    wb = openpyxl.Workbook()

    # === Sheet 1: 訂單資料 ===
    ws_orders = wb.active
    ws_orders.title = "訂單資料"

    for r in dataframe_to_rows(df_orders, index=False, header=True):
        ws_orders.append(r)

    # 格式化標題列
    header_fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    header_font = Font(bold=True, color="FFFFFF")

    for cell in ws_orders[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # 調整欄寬
    ws_orders.column_dimensions["A"].width = 12
    ws_orders.column_dimensions["B"].width = 12
    ws_orders.column_dimensions["C"].width = 12
    ws_orders.column_dimensions["D"].width = 18
    ws_orders.column_dimensions["E"].width = 10
    ws_orders.column_dimensions["F"].width = 8
    ws_orders.column_dimensions["G"].width = 10
    ws_orders.column_dimensions["H"].width = 12
    ws_orders.column_dimensions["I"].width = 10
    ws_orders.column_dimensions["J"].width = 12

    # === Sheet 2: 產品主檔 ===
    ws_products = wb.create_sheet("產品主檔")

    for r in dataframe_to_rows(df_products, index=False, header=True):
        ws_products.append(r)

    for cell in ws_products[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for col in ["A", "B", "C", "D", "E", "F"]:
        ws_products.column_dimensions[col].width = 15

    # === Sheet 3: 教學工作表 ===
    ws_tutorial = wb.create_sheet("函數教學")

    tutorial_content = [
        ["🎯 Case 01: 動態銷售報表 - 函數教學", "", "", ""],
        ["", "", "", ""],
        ["✅ 練習 1: 使用 FILTER 篩選本月訂單", "", "", ""],
        ["說明：", "篩選出 2024年11月 的所有訂單", "", ""],
        [
            "公式：",
            "=FILTER(訂單資料!A2:J1001, (訂單資料!B2:B1001>=DATE(2024,11,1))*(訂單資料!B2:B1001<=DATE(2024,11,30)))",
            "",
            "",
        ],
        ["", "", "", ""],
        ["✅ 練習 2: 使用 SORT 排序", "", "", ""],
        ["說明：", "將結果按金額降序排列", "", ""],
        ["公式：", "=SORT(FILTER(...), 7, -1)", "", ""],
        ["", "", "", ""],
        ["✅ 練習 3: 使用 UNIQUE 取得所有產品類別", "", "", ""],
        ["公式：", "=UNIQUE(訂單資料!C2:C1001)", "", ""],
        ["", "", "", ""],
        ["✅ 練習 4: 使用 XLOOKUP 查詢產品資訊", "", "", ""],
        ["說明：", "根據產品名稱查詢建議售價", "", ""],
        [
            "公式：",
            '=XLOOKUP(D2, 產品主檔!$C$2:$C$100, 產品主檔!$D$2:$D$100, "查無此產品")',
            "",
            "",
        ],
        ["", "", "", ""],
        ["✅ 練習 5: 使用 SUMIFS 計算各類別銷售額", "", "", ""],
        [
            "公式：",
            '=SUMIFS(訂單資料!$G$2:$G$1001, 訂單資料!$C$2:$C$1001, A2, 訂單資料!$J$2:$J$1001, "已完成")',
            "",
            "",
        ],
        ["", "", "", ""],
        ["✅ 練習 6: 使用 LET 簡化公式", "", "", ""],
        ["說明：", "定義變數來簡化複雜公式", "", ""],
        [
            "公式：",
            "=LET(訂單範圍, 訂單資料!A2:J1001, 月初, DATE(2024,11,1), 月底, DATE(2024,11,30), FILTER(訂單範圍, (INDEX(訂單範圍,,2)>=月初)*(INDEX(訂單範圍,,2)<=月底)))",
            "",
            "",
        ],
        ["", "", "", ""],
        ["📊 KPI 計算範例", "", "", ""],
        ["", "", "", ""],
        [
            "本月總營收：",
            '=SUMIFS(訂單資料!G:G, 訂單資料!B:B, ">=2024-11-01", 訂單資料!B:B, "<=2024-11-30", 訂單資料!J:J, "已完成")',
            "",
            "",
        ],
        [
            "本月訂單數：",
            '=COUNTIFS(訂單資料!B:B, ">=2024-11-01", 訂單資料!B:B, "<=2024-11-30", 訂單資料!J:J, "已完成")',
            "",
            "",
        ],
        ["平均客單價：", "=本月總營收/本月訂單數", "", ""],
        ["", "", "", ""],
        ["🎓 Excel → Python 對照", "", "", ""],
        ["", "", "", ""],
        ["Excel: FILTER(A:A, B:B>100)", "Python: df[df['B'] > 100]", "", ""],
        [
            "Excel: SORT(A:A, 1, -1)",
            "Python: df.sort_values('A', ascending=False)",
            "",
            "",
        ],
        ["Excel: UNIQUE(A:A)", "Python: df['A'].unique()", "", ""],
        [
            "Excel: SUMIFS(金額, 類別, '電腦')",
            "Python: df[df['類別']=='電腦']['金額'].sum()",
            "",
            "",
        ],
    ]

    for row in tutorial_content:
        ws_tutorial.append(row)

    # 格式化教學工作表
    ws_tutorial["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws_tutorial["A1"].fill = PatternFill(
        start_color="FF6B35", end_color="FF6B35", fill_type="solid"
    )
    ws_tutorial.merge_cells("A1:D1")

    ws_tutorial.column_dimensions["A"].width = 20
    ws_tutorial.column_dimensions["B"].width = 80
    ws_tutorial.column_dimensions["C"].width = 15
    ws_tutorial.column_dimensions["D"].width = 15

    # === Sheet 4: 練習區 ===
    ws_practice = wb.create_sheet("練習區")

    practice_headers = [
        ["本月訂單篩選", "", "", "", "", "", "", "", "", ""],
        ["（在這裡練習 FILTER 公式）", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", ""],
        ["Top 10 熱銷產品", "", "", "", "", "", "", "", "", ""],
        ["（在這裡練習 SORT + FILTER）", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", ""],
        ["", "", "", "", "", "", "", "", "", ""],
        ["類別銷售統計", "", "", "", "", "", "", "", "", ""],
        ["產品類別", "銷售額", "訂單數", "平均單價", "", "", "", "", "", ""],
    ]

    for row in practice_headers:
        ws_practice.append(row)

    ws_practice["A1"].font = Font(size=12, bold=True)
    ws_practice["A5"].font = Font(size=12, bold=True)
    ws_practice["A9"].font = Font(size=12, bold=True)

    for cell in ws_practice[10]:
        cell.fill = PatternFill(
            start_color="E7E6E6", end_color="E7E6E6", fill_type="solid"
        )
        cell.font = Font(bold=True)

    # 儲存檔案
    filepath = OUTPUT_DIR / "case01_dynamic_sales_report.xlsx"
    wb.save(filepath)
    print(f"✅ 已生成: {filepath}")
    print(f"   - 訂單資料: {len(df_orders)} 筆")
    print(f"   - 產品主檔: {len(df_products)} 筆")


# =============================================================================
# Case 02: 庫存儀表板
# =============================================================================
def generate_case02():
    print("\n📝 生成 Case 02: 庫存儀表板...")

    # 產品庫存資料
    categories = ["電子產品", "家電", "服飾", "食品", "日用品"]
    n_products = 200

    inventory = []
    for i in range(n_products):
        cat = np.random.choice(categories)
        current_stock = np.random.randint(0, 500)
        min_stock = np.random.randint(20, 100)
        max_stock = np.random.randint(200, 600)

        inventory.append(
            {
                "產品編號": f"SKU{i+1:05d}",
                "產品名稱": fake.word().capitalize() + " " + fake.word().capitalize(),
                "產品類別": cat,
                "當前庫存": current_stock,
                "安全庫存": min_stock,
                "最大庫存": max_stock,
                "單位成本": np.random.randint(50, 2000),
                "倉庫位置": np.random.choice(["A倉", "B倉", "C倉"]),
                "供應商": fake.company(),
                "最後進貨日": fake.date_between(start_date="-3m", end_date="today"),
            }
        )

    df_inventory = pd.DataFrame(inventory)

    # 創建 Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "庫存資料"

    for r in dataframe_to_rows(df_inventory, index=False, header=True):
        ws.append(r)

    # 格式化
    header_fill = PatternFill(
        start_color="70AD47", end_color="70AD47", fill_type="solid"
    )
    header_font = Font(bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # 調整欄寬
    for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J"]:
        ws.column_dimensions[col].width = 15

    # === 教學工作表 ===
    ws_tutorial = wb.create_sheet("函數教學")

    tutorial = [
        ["🎯 Case 02: 庫存儀表板 - 函數教學", "", ""],
        ["", "", ""],
        ["✅ 練習 1: 使用 IF 判斷庫存狀態", "", ""],
        [
            "公式：",
            '=IF(D2<E2, "缺貨", IF(D2<E2*1.5, "警告", IF(D2>F2, "過量", "正常")))',
            "",
        ],
        ["說明：", "根據當前庫存與安全庫存的關係判斷狀態", ""],
        ["", "", ""],
        ["✅ 練習 2: 使用 IFS 簡化多條件", "", ""],
        [
            "公式：",
            '=IFS(D2<E2, "缺貨", D2<E2*1.5, "警告", D2>F2, "過量", TRUE, "正常")',
            "",
        ],
        ["", "", ""],
        ["✅ 練習 3: 使用條件格式", "", ""],
        ["步驟：", "選取庫存欄位 → 常用 → 條件格式 → 資料條", ""],
        ["", "", ""],
        ["✅ 練習 4: 使用 COUNTIFS 統計各狀態產品數", "", ""],
        ["缺貨產品數：", '=COUNTIF(K:K, "缺貨")', ""],
        ["警告產品數：", '=COUNTIF(K:K, "警告")', ""],
        ["", "", ""],
        ["✅ 練習 5: 建立樞紐分析表", "", ""],
        ["步驟：", "插入 → 樞紐分析表 → 列：類別、值：庫存金額", ""],
    ]

    for row in tutorial:
        ws_tutorial.append(row)

    ws_tutorial["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws_tutorial["A1"].fill = PatternFill(
        start_color="70AD47", end_color="70AD47", fill_type="solid"
    )
    ws_tutorial.merge_cells("A1:C1")

    ws_tutorial.column_dimensions["A"].width = 20
    ws_tutorial.column_dimensions["B"].width = 80

    # 儲存
    filepath = OUTPUT_DIR / "case02_inventory_dashboard.xlsx"
    wb.save(filepath)
    print(f"✅ 已生成: {filepath}")
    print(f"   - 庫存資料: {len(df_inventory)} 筆")


# =============================================================================
# 主程式
# =============================================================================
if __name__ == "__main__":
    generate_case01()
    generate_case02()

    print("\n" + "=" * 80)
    print("🎉 所有 Excel 練習檔案生成完成！")
    print("=" * 80)
    print(f"📁 檔案位置: {OUTPUT_DIR}")
    print("\n請使用 Excel 365 或 Google Sheets 開啟這些檔案進行練習")
