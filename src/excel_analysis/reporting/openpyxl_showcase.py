from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from excel_analysis.schemas import ReportArtifact


class OpenpyxlFeatureShowcase:
    """Build a workbook dedicated to openpyxl feature demonstrations."""

    def build(self, output_path: str | Path = "reports/openpyxl_feature_showcase.xlsx") -> ReportArtifact:
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        wb.remove(wb.active)
        self._register_styles(wb)
        self._capabilities(wb)
        self._formula_demo(wb)
        self._named_styles_demo(wb)
        self._comments_hyperlinks_demo(wb)
        self._data_validation_demo(wb)
        self._protection_demo(wb)
        self._conditional_formatting_demo(wb)
        self._charts_demo(wb)
        wb.save(output)
        return ReportArtifact(output, "excel", "Standalone openpyxl feature showcase workbook")

    def _register_styles(self, wb: Workbook) -> None:
        styles = {
            "showcase_title": {
                "font": Font(bold=True, size=16, color="FFFFFF"),
                "fill": PatternFill("solid", fgColor="1F4E78"),
                "alignment": Alignment(horizontal="center"),
            },
            "showcase_header": {
                "font": Font(bold=True, color="FFFFFF"),
                "fill": PatternFill("solid", fgColor="4472C4"),
                "alignment": Alignment(horizontal="center"),
                "border": Border(bottom=Side(style="thin", color="808080")),
            },
            "showcase_currency": {"number_format": '$#,##0.00'},
            "showcase_percent": {"number_format": "0.0%"},
        }
        for name, attrs in styles.items():
            if name in wb.named_styles:
                continue
            style = NamedStyle(name=name)
            for attr, value in attrs.items():
                setattr(style, attr, value)
            wb.add_named_style(style)

    def _capabilities(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Capabilities")
        ws.append(["Feature", "Worksheet", "Openpyxl API"])
        rows = [
            ("Formula demo", "Formula Demo", "cell.value = '=SUM(...)'"),
            ("Named styles", "Named Styles", "NamedStyle, Font, Fill, Border"),
            ("Comments and hyperlinks", "Comments Links", "Comment, cell.hyperlink"),
            ("Data validation", "Data Validation", "DataValidation"),
            ("Protection", "Protection", "Worksheet.protection"),
            ("Conditional formatting", "Conditional Formatting", "ColorScaleRule, DataBarRule, CellIsRule"),
            ("Charts", "Charts", "BarChart, LineChart, PieChart"),
        ]
        for row in rows:
            ws.append(row)
        self._table(ws, "CapabilitiesTable")
        self._finish(ws)

    def _formula_demo(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Formula Demo")
        ws.append(["Month", "Revenue", "Cost", "Profit", "Margin"])
        values = [
            ("Jan", 125000, 78000),
            ("Feb", 142000, 83000),
            ("Mar", 138500, 80500),
            ("Apr", 155200, 91000),
        ]
        for row_idx, (month, revenue, cost) in enumerate(values, 2):
            ws.cell(row_idx, 1, month)
            ws.cell(row_idx, 2, revenue)
            ws.cell(row_idx, 3, cost)
            ws.cell(row_idx, 4, f"=B{row_idx}-C{row_idx}")
            ws.cell(row_idx, 5, f"=D{row_idx}/B{row_idx}")
        ws.append(["Total", "=SUM(B2:B5)", "=SUM(C2:C5)", "=SUM(D2:D5)", "=D6/B6"])
        ws["A8"] = "Formula cells are intentionally stored as formulas for Excel recalculation."
        self._style_header(ws)
        self._finish(ws)

    def _named_styles_demo(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Named Styles")
        ws.merge_cells("A1:D1")
        ws["A1"] = "NamedStyle Portfolio Formatting"
        ws["A1"].style = "showcase_title"
        ws.append(["Metric", "Value", "Format", "Notes"])
        rows = [
            ("Revenue", 125000, "currency", "Currency named style"),
            ("Gross Margin", 0.382, "percent", "Percent named style"),
            ("Status", "On Track", "header", "Header style used across reports"),
        ]
        for row in rows:
            ws.append(row)
        ws["B3"].style = "showcase_currency"
        ws["B4"].style = "showcase_percent"
        self._style_header(ws, row=2)
        self._finish(ws)

    def _comments_hyperlinks_demo(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Comments Links")
        ws.append(["Resource", "Link", "Comment"])
        ws.append(["Openpyxl Docs", "https://openpyxl.readthedocs.io/", "Official documentation"])
        ws.append(["Generated Report", "#Capabilities!A1", "Internal workbook navigation"])
        ws["B2"].hyperlink = "https://openpyxl.readthedocs.io/"
        ws["B2"].style = "Hyperlink"
        ws["B3"].hyperlink = "#Capabilities!A1"
        ws["B3"].style = "Hyperlink"
        ws["A1"].comment = Comment("This sheet demonstrates comments and external/internal hyperlinks.", "excel_analysis")
        self._style_header(ws)
        self._finish(ws)

    def _data_validation_demo(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Data Validation")
        ws.append(["Order ID", "Status", "Priority", "Approval"])
        for idx in range(2, 12):
            ws.cell(idx, 1, f"ORD-{idx:03d}")
        status = DataValidation(type="list", formula1='"New,Processing,Shipped,Cancelled"', allow_blank=False)
        priority = DataValidation(type="list", formula1='"Low,Medium,High"', allow_blank=False)
        approval = DataValidation(type="whole", operator="between", formula1="0", formula2="1")
        ws.add_data_validation(status)
        ws.add_data_validation(priority)
        ws.add_data_validation(approval)
        status.add("B2:B11")
        priority.add("C2:C11")
        approval.add("D2:D11")
        self._style_header(ws)
        self._finish(ws)

    def _protection_demo(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Protection")
        ws.append(["Area", "Editable", "Notes"])
        ws.append(["Forecast assumption", "Yes", "Input cells unlocked in a real business template"])
        ws.append(["Calculated total", "No", "Protected output area"])
        ws["A1"].comment = Comment("Sheet protection is enabled for this demo.", "excel_analysis")
        ws.protection.sheet = True
        ws.protection.password = "demo"
        self._style_header(ws)
        self._finish(ws)

    def _conditional_formatting_demo(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Conditional Formatting")
        ws.append(["SKU", "Inventory Health", "Revenue", "Margin"])
        rows = [
            ("SKU-001", 0.92, 120000, 0.34),
            ("SKU-002", 0.55, 54000, 0.19),
            ("SKU-003", 0.24, 97000, 0.42),
            ("SKU-004", 0.78, 73000, 0.28),
            ("SKU-005", 0.12, 23000, 0.08),
        ]
        for row in rows:
            ws.append(row)
        ws.conditional_formatting.add("B2:B6", ColorScaleRule(start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"))
        ws.conditional_formatting.add("C2:C6", DataBarRule(start_type="min", end_type="max", color="4472C4"))
        ws.conditional_formatting.add("D2:D6", CellIsRule(operator="lessThan", formula=["0.2"], fill=PatternFill("solid", fgColor="FFC7CE")))
        self._style_header(ws)
        self._finish(ws)

    def _charts_demo(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Charts")
        ws.append(["Month", "Revenue", "Orders", "Profit"])
        rows = [
            ("Jan", 125000, 890, 47000),
            ("Feb", 142000, 940, 59000),
            ("Mar", 138500, 910, 58000),
            ("Apr", 155200, 1005, 64200),
            ("May", 162500, 1050, 68000),
        ]
        for row in rows:
            ws.append(row)
        self._style_header(ws)
        bar = BarChart()
        bar.title = "Revenue by Month"
        bar.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
        ws.add_chart(bar, "F2")
        line = LineChart()
        line.title = "Orders Trend"
        line.add_data(Reference(ws, min_col=3, min_row=1, max_row=6), titles_from_data=True)
        line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
        ws.add_chart(line, "F18")
        pie = PieChart()
        pie.title = "Profit Mix"
        pie.add_data(Reference(ws, min_col=4, min_row=1, max_row=6), titles_from_data=True)
        pie.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
        ws.add_chart(pie, "N2")
        self._finish(ws)

    def _style_header(self, ws, row: int = 1) -> None:
        for cell in ws[row]:
            cell.style = "showcase_header"

    def _table(self, ws, name: str) -> None:
        table = Table(displayName=name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
        ws.add_table(table)
        self._style_header(ws)

    def _finish(self, ws) -> None:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, column_cells in enumerate(ws.columns, 1):
            width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:100])
            ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 12), 36)
