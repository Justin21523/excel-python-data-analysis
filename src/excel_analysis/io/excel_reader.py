from __future__ import annotations

from pathlib import Path
from typing import BinaryIO

import pandas as pd
from openpyxl import load_workbook


class ExcelReader:
    """Excel workbook reader with portfolio-friendly inspection features."""

    def read(self, source: str | Path | BinaryIO, sheet_name: str | int | None = None) -> dict[str, pd.DataFrame]:
        sheets = pd.read_excel(source, sheet_name=sheet_name, engine="openpyxl")
        if isinstance(sheets, pd.DataFrame):
            key = str(sheet_name if sheet_name is not None else "Sheet1")
            return {key: self._clean_frame(sheets)}
        return {name: self._clean_frame(df) for name, df in sheets.items()}

    def inspect_workbook(self, path: str | Path) -> dict[str, dict[str, object]]:
        workbook = load_workbook(path, data_only=False)
        inspection: dict[str, dict[str, object]] = {}
        for ws in workbook.worksheets:
            merged_ranges = [str(item) for item in ws.merged_cells.ranges]
            formula_cells = [
                cell.coordinate
                for row in ws.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            inspection[ws.title] = {
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "merged_ranges": merged_ranges,
                "formula_cells": formula_cells,
            }
        return inspection

    def _clean_frame(self, df: pd.DataFrame) -> pd.DataFrame:
        cleaned = df.dropna(how="all").dropna(axis=1, how="all").copy()
        cleaned.columns = [str(col).strip() for col in cleaned.columns]
        return cleaned

